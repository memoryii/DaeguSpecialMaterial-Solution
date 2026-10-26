"""Reproducible MES LOT analysis used by the seven-step company notebook.

Input filenames, exact worksheet names and source-specific notes stay in the
private analysis_config.json. The public module contains no company records.
"""

from pathlib import Path
from datetime import datetime, timezone
from collections import Counter
from importlib import metadata
from html import escape
import hashlib
import json
import os
import uuid

import pandas as pd
import openpyxl
import matplotlib.pyplot as plt
from matplotlib import font_manager
from IPython.display import display, Markdown, HTML


def _sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _records(frame):
    return json.loads(
        frame.to_json(orient="records", force_ascii=False, date_format="iso")
    )


def _details(title, *items):
    """Keep full diagnostic tables available without expanding the default view."""
    parts = []
    for item in items:
        if isinstance(item, pd.DataFrame):
            parts.append(item.to_html(index=False, escape=True))
        elif isinstance(item, (list, tuple)):
            parts.append(
                "<ul>"
                + "".join("<li>" + escape(str(line)) + "</li>" for line in item)
                + "</ul>"
            )
        else:
            parts.append("<p>" + escape(str(item)) + "</p>")
    display(
        HTML(
            "<details><summary>"
            + escape(title)
            + "</summary>"
            + "".join(parts)
            + "</details>"
        )
    )


class MESAnalysis:
    """State for one ordered run; instantiate again before rerunning all cells."""

    def __init__(self):
        work = Path.cwd()
        default_input = (
            work / "data"
            if (work / "data").is_dir()
            else Path.home() / "daegu_mes_private" / "inputs"
        )
        default_output = (
            work / "outputs"
            if (work / "outputs").is_dir()
            else Path.home() / "daegu_mes_private" / "outputs"
        )
        self.input_dir = (
            Path(os.environ.get("MES_DATA_DIR", default_input)).expanduser().resolve()
        )
        self.output_dir = (
            Path(os.environ.get("MES_OUTPUT_DIR", default_output))
            .expanduser()
            .resolve()
        )
        source_root = Path(__file__).resolve().parents[1]
        self.public_root = (
            Path(os.environ.get("MES_REPOSITORY_DIR", source_root))
            .expanduser()
            .resolve()
        )
        if self.input_dir.is_relative_to(
            self.public_root
        ) or self.output_dir.is_relative_to(self.public_root):
            raise ValueError(
                "실제 입력과 결과는 공개 저장소 바깥의 private 경로를 사용하세요."
            )
        if self.output_dir.is_relative_to(
            self.input_dir
        ) or self.input_dir.is_relative_to(self.output_dir):
            raise ValueError("입력과 출력 경로는 서로 겹치지 않아야 합니다.")
        config_path = self._input_file("analysis_config.json")
        if not config_path.is_file():
            raise FileNotFoundError("data 폴더에 analysis_config.json이 필요합니다.")
        self.config = json.loads(config_path.read_text(encoding="utf-8"))
        self.files = self.config["files"]
        self.sheets = self.config["sheets"]
        self.processes = self.config["process_names"]
        self.original = self.files["original"]
        self.supplement = self.files["supplement"]
        self.run_dir = self.output_dir / (
            "run_"
            + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
            + "_"
            + uuid.uuid4().hex[:6]
        )
        required = [config_path] + [
            self._input_file(self.files[key])
            for key in [
                "original",
                "supplement",
                "detail_csv",
                "defect_csv",
                "field_mapping_csv",
            ]
        ]
        required += [self._input_file(name) for name in self.files["source_images"]]
        missing = [str(path) for path in required if not path.is_file()]
        if missing:
            raise FileNotFoundError("필수 입력 누락: " + "; ".join(missing))
        self.before = self._snapshot()
        self.environment = {
            "python": os.sys.version.split()[0],
            **{
                p: metadata.version(p)
                for p in ["pandas", "openpyxl", "matplotlib", "IPython"]
            },
        }

    def _input_file(self, relative_path):
        candidate = (self.input_dir / relative_path).resolve()
        if not candidate.is_relative_to(self.input_dir):
            raise ValueError("입력 설정은 data 폴더 내부의 파일만 지정할 수 있습니다.")
        return candidate

    def _snapshot(self):
        return {
            str(path.relative_to(self.input_dir)): _sha256(path)
            for path in sorted(self.input_dir.rglob("*"))
            if path.is_file()
            and ".ipynb_checkpoints" not in path.relative_to(self.input_dir).parts
        }

    def _require(self, *names):
        if any(not hasattr(self, name) for name in names):
            raise RuntimeError(
                "앞 단계부터 순서대로 실행하세요. Run → Run All Cells로 전체 실행할 수 있습니다."
            )

    def _read_exact_table(self, sheet):
        wb = openpyxl.load_workbook(
            self._input_file(self.original), data_only=True, read_only=True
        )
        try:
            ws = wb[sheet]
            values = list(ws.values)
            header = list(values[0])
            if len(set(header)) != len(header):
                raise ValueError("중복 헤더: " + sheet)
            # object preserves integer sequence keys when footer rows contain blanks.
            frame = pd.DataFrame(values[1:], columns=header, dtype=object)
            frame["출처행"] = range(2, len(frame) + 2)
            frame["출처파일"] = self.original
            frame["출처시트"] = sheet
            return frame
        finally:
            wb.close()

    def load_data(self):
        "원본 시트·컬럼 인벤토리와 전사 CSV를 읽고 원문 키를 보존합니다."
        inventory = []
        all_cell_data = []
        column_inventory = []
        for file_role in ["original", "supplement"]:
            filename = self.files[file_role]
            wb = openpyxl.load_workbook(
                self._input_file(filename), data_only=True, read_only=False
            )
            for ws in wb:
                populated = [
                    (i, list(row))
                    for i, row in enumerate(ws.iter_rows(values_only=True), 1)
                    if any((value is not None for value in row))
                ]
                rule = self.config["sheet_rules"][file_role][ws.title]
                inventory.append(
                    {
                        "파일": filename,
                        "시트": ws.title,
                        "사용범위행수": ws.max_row,
                        "사용범위열수": ws.max_column,
                        "비어있지않은행수": len(populated),
                        "삽입이미지수": len(ws._images),
                        "분석대상셀레코드수": rule["analysis_rows"],
                        "레코드의미": rule["record_meaning"],
                        "원문영역": rule["source_region"],
                        "전체모집단표본수": "확정 불가",
                    }
                )
                header_ranges = rule.get(
                    "header_ranges",
                    [f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}1"],
                )
                for header_range in header_ranges:
                    for row in ws[header_range]:
                        for cell in row:
                            if cell.value is not None:
                                column_inventory.append(
                                    {
                                        "파일": filename,
                                        "시트": ws.title,
                                        "헤더영역": header_range,
                                        "헤더셀": cell.coordinate,
                                        "원문컬럼명": str(cell.value),
                                    }
                                )
                for row_no, row in populated:
                    for col_no, value in enumerate(row, 1):
                        if value is not None:
                            all_cell_data.append(
                                {
                                    "파일": filename,
                                    "시트": ws.title,
                                    "셀": ws.cell(row_no, col_no).coordinate,
                                    "값": (
                                        value.isoformat(sep=" ")
                                        if isinstance(value, datetime)
                                        else value
                                    ),
                                }
                            )
            wb.close()
        self.sheet_inventory = pd.DataFrame(inventory)
        self.source_cells = pd.DataFrame(all_cell_data)
        self.column_inventory = pd.DataFrame(column_inventory)
        self.summary = self._read_exact_table(self.sheets["summary"])
        self.material_all = self._read_exact_table(self.sheets["material"])
        self.materials = self.material_all.loc[
            self.material_all["지시번호"].notna()
        ].copy()
        self.material_footer = self.material_all.loc[
            self.material_all["지시번호"].isna()
        ].copy()
        self.printing_passes = self._read_exact_table(self.sheets["printing"])
        self.inventory_stock = self._read_exact_table(self.sheets["stock"])
        self.detail = pd.read_csv(
            self._input_file(self.files["detail_csv"]),
            dtype="string",
            keep_default_na=False,
        )
        self.defects = pd.read_csv(
            self._input_file(self.files["defect_csv"]),
            dtype="string",
            keep_default_na=False,
        )
        self.field_mapping = pd.read_csv(
            self._input_file(self.files["field_mapping_csv"]),
            dtype="string",
            keep_default_na=False,
        )
        for frame in [
            self.summary,
            self.materials,
            self.detail,
            self.defects,
            self.printing_passes,
        ]:
            for column in ["지시번호", "작업순서", "실적순번", "LOT NO", "품목코드"]:
                if column in frame:
                    frame[column] = frame[column].astype("string")
        for frame, columns in [
            (self.summary, ["지시수량", "양품수량", "불량수량"]),
            (self.detail, ["지시수량", "실적수량", "양품수량", "불량수량"]),
            (self.materials, ["소요량"]),
            (self.defects, ["불량수량"]),
        ]:
            for column in columns:
                frame[column] = pd.to_numeric(frame[column], errors="raise")
        print(
            f"입력 확인: 원본 {len(self.sheet_inventory['파일'].unique())}개 파일 · {len(self.sheet_inventory)}개 시트 · 상세 LOT {len(self.detail)}행"
        )
        _details(
            "파일·시트·컬럼 상세 보기",
            self.sheet_inventory,
            self.column_inventory,
            self.field_mapping,
        )

    def check_data(self):
        "누락·중복·기록 기간과 제공된 키의 매칭 관계를 점검합니다."
        self._require("detail", "sheet_inventory")
        self.key = ["지시번호", "작업순서", "실적순번", "LOT NO"]

        def quality_row(name, frame, keys):
            null = frame[keys].isna().any(axis=1) | frame[keys].eq("").any(axis=1)
            valid = frame.loc[~null, keys]
            return {
                "테이블": name,
                "관측행수": len(frame),
                "검토키": " + ".join(keys),
                "키누락행수": int(null.sum()),
                "고유키수": len(valid.drop_duplicates()),
                "중복그룹소속행수": int(valid.duplicated(keep=False).sum()),
                "추가중복행수": int(valid.duplicated(keep="first").sum()),
                "범위": "제공된 기록 한정; 모집단 품질률 아님",
            }

        self.quality = pd.DataFrame(
            [
                quality_row("상세 LOT", self.detail, self.key),
                quality_row("공정 요약", self.summary, ["지시번호", "작업순서"]),
                quality_row(
                    "자재 소요량",
                    self.materials,
                    ["지시번호", "작업순서", "실적순번", "소요량순번"],
                ),
                quality_row(
                    "불량 유형",
                    self.defects,
                    ["지시번호", "작업순서", "실적순번", "LOT NO", "불량코드"],
                ),
            ]
        )
        self.summary_time_equal = bool(
            (
                pd.to_datetime(self.summary["시작시간"])
                == pd.to_datetime(self.summary["종료시간"])
            ).all()
        )
        self.summary_same_timestamp = (
            int(
                pd.concat(
                    [self.summary["시작시간"], self.summary["종료시간"]]
                ).nunique()
            )
            == 1
        )
        self.detail["화면시작_파싱"] = pd.to_datetime(
            self.detail["시작시간"], format="%Y-%m-%d %H:%M", errors="raise"
        )
        self.detail["화면종료_파싱"] = pd.to_datetime(
            self.detail["종료시간"], format="%Y-%m-%d %H:%M", errors="raise"
        )
        self.detail["기록경과시간_분"] = (
            self.detail["화면종료_파싱"] - self.detail["화면시작_파싱"]
        ).dt.total_seconds() / 60
        if (self.detail["기록경과시간_분"] < 0).any():
            raise ValueError("화면 시각의 종료가 시작보다 빠릅니다.")
        self.period = pd.DataFrame(
            [
                {
                    "대상": f"상세 LOT 화면 {len(self.detail)}행",
                    "시간컬럼": "시작시간/종료시간",
                    "시작": str(self.detail["화면시작_파싱"].min()),
                    "종료": str(self.detail["화면종료_파싱"].max()),
                    "해석": "관측된 화면 범위; 전체 생산기간 아님",
                },
                {
                    "대상": f"{self.sheets['printing']} ({len(self.printing_passes)}행)",
                    "시간컬럼": "실적시작/실적종료",
                    "시작": str(self.printing_passes["실적시작"].min()),
                    "종료": str(self.printing_passes["실적종료"].max()),
                    "해석": "인쇄 상세 기록 범위; 초/밀리초 원본 유지",
                },
                {
                    "대상": f"{self.sheets['summary']} ({len(self.summary)}행)",
                    "시간컬럼": "시작시간/종료시간",
                    "시작": str(self.summary["시작시간"].min()),
                    "종료": str(self.summary["종료시간"].max()),
                    "해석": (
                        "모든 시작/종료가 동일: 실제 소요/대기시간 분석 제외"
                        if self.summary_same_timestamp
                        else "시각의 업무 의미 확인 후 사용"
                    ),
                },
            ]
        )
        self.missing = pd.DataFrame(
            [
                {
                    "항목": "공정요약 설비명 미기재",
                    "누락수": int(self.summary["설비명"].isna().sum()),
                    "분모": len(self.summary),
                    "의미": "요약 컬럼의 미기재; 상세 설비 부재로 단정 불가",
                },
                {
                    "항목": "상세LOT 최종검사 불량귀속",
                    "누락수": None,
                    "분모": None,
                    "의미": "연결 데이터 미제공; 최종 불량을 상류/원료LOT에 임의 배분하지 않음",
                },
                {
                    "항목": "재고 오더번호 미기재",
                    "누락수": int(self.inventory_stock["오더번호"].isna().sum()),
                    "분모": len(self.inventory_stock),
                    "의미": "재고 스냅샷 한정; 본 작업지시 실적과 분리",
                },
            ]
        )

        def audit_link(name, left, right, keys, meaning):

            def valid_counts(frame):
                ok = frame[keys].notna().all(axis=1) & ~frame[keys].eq("").any(axis=1)
                return (
                    ok,
                    Counter(frame.loc[ok, keys].itertuples(index=False, name=None)),
                )

            lok, lc = valid_counts(left)
            rok, rc = valid_counts(right)
            common = set(lc) & set(rc)
            matched = sum((lc[k] for k in common))
            rel = Counter(
                (
                    ("1" if lc[k] == 1 else "N") + ":" + ("1" if rc[k] == 1 else "N")
                    for k in common
                )
            )
            return {
                "연결": name,
                "키": " + ".join(keys),
                "왼쪽행수": len(left),
                "키누락행수": int((~lok).sum()),
                "매칭행수": matched,
                "행기준매칭률_pct": 100 * matched / len(left) if len(left) else None,
                "유효키고유수": len(lc),
                "매칭고유키수": len(common),
                "고유키매칭률_pct": 100 * len(common) / len(lc) if lc else None,
                "1대1키수": rel["1:1"],
                "1대N키수": rel["1:N"],
                "N대1키수": rel["N:1"],
                "N대N키수": rel["N:N"],
                "의미": meaning,
                "범위": "현재 제공행 한정; 전체 공정 이력 복원율 아님",
            }

        self.matching = pd.DataFrame(
            [
                audit_link(
                    "상세LOT → 공정그룹",
                    self.detail,
                    self.summary,
                    ["지시번호", "작업순서"],
                    "상세 실적과 공정 요약의 집계 구조 연결",
                ),
                audit_link(
                    "상세LOT → 요약LOT 값 대조",
                    self.detail,
                    self.summary,
                    ["지시번호", "LOT NO"],
                    "요약에 별도 LOT행이 없을 수 있음; 집계 차이를 오류로 단정하지 않음",
                ),
                audit_link(
                    "자재행 → 보충 화면의 작업실적내역",
                    self.materials,
                    self.detail,
                    ["지시번호", "작업순서", "실적순번"],
                    "자재행과 작업실적내역의 관계; 제품/불량의 원료 귀속은 별도 정보 필요",
                ),
                audit_link(
                    "불량유형 → 보충 화면의 작업실적내역",
                    self.defects,
                    self.detail,
                    self.key,
                    "해당 실적의 불량 유형 설명이며 같은 수량을 추가 불량으로 합산하지 않음",
                ),
            ]
        )
        display(self.quality[["테이블", "관측행수", "키누락행수", "추가중복행수"]])
        display(self.matching[["연결", "왼쪽행수", "매칭행수", "행기준매칭률_pct"]])
        _details(
            "누락·기록 기간·연결 상세 보기",
            self.quality,
            self.missing,
            self.period,
            self.matching,
        )

    def summarize(self):
        "기존 수량 검산·표시 불량비율·자재 구성비를 계산합니다."
        self._require("matching")
        self.pressure = self.detail.loc[
            self.detail["공정명"].eq(self.processes["pressure"])
        ]
        self.printing = self.detail.loc[
            self.detail["공정명"].eq(self.processes["printing"])
        ]
        self.final = self.summary.loc[
            self.summary["공정명"].eq(self.processes["inspection"])
        ].iloc[0]
        self.inspection_total = int(self.final["양품수량"] + self.final["불량수량"])
        self.material_sum = float(self.materials["소요량"].sum())
        self.material_view = self.materials[
            [
                "지시번호",
                "작업순서",
                "실적순번",
                "소요량순번",
                "품목코드",
                "품목명",
                "LOT NO",
                "소요량",
            ]
        ].copy()
        self.material_view["소요량구성비_pct"] = (
            100 * self.material_view["소요량"] / self.material_sum
        )
        self.material_view["소요량단위"] = "원본 미기재"
        checks = [
            {
                "검산": "대상 공정 상세 양품 합 ↔ 공정 요약 양품",
                "왼쪽": int(self.pressure["양품수량"].sum()),
                "오른쪽": int(
                    self.summary.loc[
                        self.summary["공정명"].eq(self.processes["pressure"]),
                        "양품수량",
                    ].iloc[0]
                ),
                "의미": "해당 상세LOT의 집계 일치",
            },
            {
                "검산": "선택 공정 상세 불량 ↔ 불량유형 수량",
                "왼쪽": int(self.printing["불량수량"].sum()),
                "오른쪽": int(self.defects["불량수량"].sum()),
                "의미": "같은 불량수량의 유형 설명; 합산 금지",
            },
            {
                "검산": "자재 소요량 합 ↔ 원본 합계행",
                "왼쪽": self.material_sum,
                "오른쪽": float(self.material_footer["소요량"].iloc[0]),
                "의미": "합계행은 자재 표본에서 제외",
            },
            {
                "검산": "공정별 불량 합 ↔ 지시수량−최종양품",
                "왼쪽": int(self.summary["불량수량"].sum()),
                "오른쪽": int(
                    self.summary.iloc[0]["지시수량"] - self.final["양품수량"]
                ),
                "의미": "수량 산술검산; 계보/중복계수 입증 아님",
            },
        ]
        self.checks = pd.DataFrame(checks)
        self.checks["일치"] = self.checks["왼쪽"].eq(self.checks["오른쪽"])
        self.rates = pd.DataFrame(
            [
                {
                    "지표": "선택 공정 관측LOT 불량비율",
                    "분자": int(self.printing["불량수량"].sum()),
                    "분모": int(self.printing["실적수량"].sum()),
                    "pct": 100
                    * self.printing["불량수량"].sum()
                    / self.printing["실적수량"].sum(),
                    "한계": "제공된 LOT 관측 한정",
                },
                {
                    "지표": "최종 검사 표시 불량비율",
                    "분자": int(self.final["불량수량"]),
                    "분모": self.inspection_total,
                    "pct": 100 * self.final["불량수량"] / self.inspection_total,
                    "한계": "양품+불량을 표시 분모로 사용; 개별검사/재검 원본 없음",
                },
                {
                    "지표": "조건부 누적 손실 산술",
                    "분자": int(self.summary["불량수량"].sum()),
                    "분모": int(self.summary.iloc[0]["지시수량"]),
                    "pct": 100
                    * self.summary["불량수량"].sum()
                    / self.summary.iloc[0]["지시수량"],
                    "한계": "추가투입·재작업·중복계수 없다는 조건 확인 필요",
                },
            ]
        )
        assert self.checks["일치"].all()
        display(self.detail[["공정명", "LOT NO", "실적수량", "양품수량", "불량수량"]])
        display(self.rates)
        _details(
            "수량 검산·자재·불량유형 상세 보기",
            self.checks,
            self.material_view,
            self.defects,
        )

    def plot(self):
        "관측 LOT 수량·공정 요약 불량·기록 경과시간의 기존 세 그림을 표시합니다."
        self._require("checks", "period")
        font_path = os.environ.get("MES_KOREAN_FONT", "")
        self.font_name = None
        if font_path and Path(font_path).is_file():
            font_manager.fontManager.addfont(font_path)
            self.font_name = font_manager.FontProperties(fname=font_path).get_name()
        if not self.font_name:
            installed = {font.name for font in font_manager.fontManager.ttflist}
            self.font_name = next(
                (
                    name
                    for name in [
                        "Noto Sans CJK KR",
                        "Malgun Gothic",
                        "AppleGothic",
                        "NanumGothic",
                    ]
                    if name in installed
                ),
                None,
            )
        self.korean = self.font_name is not None
        if self.korean:
            plt.rcParams["font.family"] = self.font_name
        plt.rcParams["axes.unicode_minus"] = False
        self.figures = {}
        fig, ax = plt.subplots(figsize=(10, 4))
        labels = self.detail["LOT NO"].tolist()
        ax.bar(
            labels,
            self.detail["양품수량"],
            label="양품" if self.korean else "Good",
            color="#347d91",
        )
        ax.bar(
            labels,
            self.detail["불량수량"],
            bottom=self.detail["양품수량"],
            label="불량" if self.korean else "Defect",
            color="#bc5545",
        )
        quantity_max = max(float(self.detail["실적수량"].max()), 1)
        for i, row in self.detail.iterrows():
            ax.text(
                i,
                row["실적수량"] + quantity_max * 0.01,
                f"{row['실적수량']:,} / {row['불량수량']:,}",
                ha="center",
            )
        ax.set_ylim(0, quantity_max * 1.15)
        ax.legend(frameon=False)
        ax.set_title(
            f"확인된 상세 LOT {len(self.detail)}행: 전체 / 불량 수량 (서로 다른 공정, 합산 금지)"
            if self.korean
            else f"{len(self.detail)} observed LOT records: total / defect (do not sum across processes)"
        )
        ax.set_ylabel("수량" if self.korean else "Count")
        fig.tight_layout()
        display(fig)
        self.figures["01_detail_lot_quantities.png"] = fig
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(10, 4))
        process_labels = (
            self.summary["공정명"].tolist()
            if self.korean
            else [
                self.config["process_labels_en"].get(name, f"Process {i + 1}")
                for i, name in enumerate(self.summary["공정명"])
            ]
        )
        bars = ax.bar(process_labels, self.summary["불량수량"], color="#bc5545")
        ax.bar_label(bars, padding=3)
        ax.set_ylim(0, max(float(self.summary["불량수량"].max()), 1) * 1.2)
        ax.set_ylabel("불량수량" if self.korean else "Defect count")
        ax.set_title(
            "원본 공정 요약의 불량수량 — 최종 불량의 상세LOT 귀속은 별도 확인 필요"
            if self.korean
            else "Process-summary defects; final defects not allocated to detail LOTs"
        )
        fig.tight_layout()
        display(fig)
        self.figures["02_process_summary_defects.png"] = fig
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(10, 4))
        bars = ax.bar(labels, self.detail["기록경과시간_분"] / 60, color="#347d91")
        elapsed_max = max(float(self.detail["기록경과시간_분"].max() / 60), 1)
        for bar, minutes in zip(bars, self.detail["기록경과시간_분"]):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + elapsed_max * 0.01,
                f"{int(minutes // 60)}h {int(minutes % 60):02d}m",
                ha="center",
            )
        ax.set_ylim(0, elapsed_max * 1.15)
        ax.set_ylabel("시간" if self.korean else "Hours")
        ax.set_title(
            "화면 시작~종료의 달력상 경과시간 — 가동시간·공정 대기시간 아님"
            if self.korean
            else "Recorded elapsed time; not active runtime or inter-process waiting"
        )
        fig.tight_layout()
        display(fig)
        self.figures["03_recorded_lot_elapsed_time.png"] = fig
        plt.close(fig)

    def show_limits(self):
        "현재 입력의 진행 범위와 모든 추가 확인사항을 보존합니다."
        self._require("checks")
        self.status = pd.DataFrame(
            self.config["scope_status"],
            columns=["요청기능", "상태", "이번 결과", "추가 필요"],
        )
        self.limitations = self.config["limitations"]
        display(self.status[["요청기능", "상태"]])
        _details("추가 필요자료·해석 시 주의사항 보기", self.status, self.limitations)

    def save_results(self):
        "전체 표·그림·보고서를 저장하고 입력 원본의 해시를 검증합니다."
        self._require("figures", "status", "matching")
        self.run_dir.mkdir(parents=True, exist_ok=False)
        self.tables = {
            "sheet_inventory.csv": self.sheet_inventory,
            "source_cells.csv": self.source_cells,
            "detail_lot_records.csv": self.detail,
            "process_summary_records.csv": self.summary,
            "material_records.csv": self.materials,
            "material_footer.csv": self.material_footer,
            "material_quantity_composition.csv": self.material_view,
            "defect_records.csv": self.defects,
            "field_mapping.csv": self.field_mapping,
            "column_inventory.csv": self.column_inventory,
            "quality_checks.csv": self.quality,
            "missing_data_checks.csv": self.missing,
            "recorded_periods.csv": self.period,
            "matching_audit.csv": self.matching,
            "quantity_reconciliation.csv": self.checks,
            "qualified_rates.csv": self.rates,
            "scope_status.csv": self.status,
        }
        for name, frame in self.tables.items():
            frame.to_csv(self.run_dir / name, index=False, encoding="utf-8-sig")
        for name, fig in self.figures.items():
            fig.savefig(self.run_dir / name, dpi=160, bbox_inches="tight")
        self.after = self._snapshot()
        self.originals_unchanged = self.before == self.after
        self.report = {
            "product": {
                column: values[0] if len(values) == 1 else values
                for column in ["품목코드", "품목명", "지시번호"]
                for values in [self.summary[column].dropna().unique().tolist()]
            },
            "detail_lot_records": _records(self.detail),
            "process_summary_records": _records(self.summary),
            "material_records": _records(self.materials),
            "defect_records": _records(self.defects),
            "field_mapping": _records(self.field_mapping),
            "matching_audit": _records(self.matching),
            "inventory": _records(self.sheet_inventory),
            "column_inventory": _records(self.column_inventory),
            "quality_checks": _records(self.quality),
            "missing_data_checks": _records(self.missing),
            "recorded_periods": _records(self.period),
            "quantity_reconciliation": _records(self.checks),
            "qualified_rates": _records(self.rates),
            "scope_status": _records(self.status),
            "metrics": {
                "order_count": self.detail["지시번호"].nunique(),
                "detail_lot_count": len(self.detail),
                "observed_process_group_count": self.detail["작업순서"].nunique(),
                "summary_process_group_count": len(self.summary),
                "printing_good": int(self.printing["양품수량"].sum()),
                "printing_defects": int(self.printing["불량수량"].sum()),
                "pressure_detail_good_sum": int(self.pressure["양품수량"].sum()),
                "final_good": int(self.final["양품수량"]),
                "final_defects": int(self.final["불량수량"]),
                "displayed_final_inspection_total": self.inspection_total,
                "displayed_final_defect_rate_pct": 100
                * self.final["불량수량"]
                / self.inspection_total,
                "cumulative_displayed_defects": int(self.summary["불량수량"].sum()),
                "conditional_cumulative_loss_pct": 100
                * self.summary["불량수량"].sum()
                / self.summary.iloc[0]["지시수량"],
            },
            "limitations": self.limitations,
            "execution_status": "ALL_CELLS_EXECUTED",
            "run_directory": str(self.run_dir.relative_to(self.output_dir)),
            "input_sha256_before": self.before,
            "input_sha256_after": self.after,
            "originals_unchanged": self.originals_unchanged,
            "environment": self.environment,
            "chart_font": self.font_name if self.korean else "English fallback",
        }
        (self.run_dir / "report_data.json").write_text(
            json.dumps(self.report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        (self.output_dir / "report_data.json").write_text(
            json.dumps(self.report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        if not self.originals_unchanged:
            raise RuntimeError(
                "입력 SHA256이 실행 전후 달라졌습니다. report_data.json을 확인하세요."
            )
        relative = self.run_dir.relative_to(self.output_dir).as_posix()
        display(
            Markdown(
                f"**분석 완료** · [결과 폴더 열기](outputs/{relative}/) · [입력 자료 열기](data/)"
            )
        )
        print(
            f"작업지시 {self.detail['지시번호'].nunique()}건 · 상세 LOT {len(self.detail)}행 · 입력 원본 보존 확인"
        )
        return self.run_dir
